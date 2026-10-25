"""
convert_report_to_docx.py — Convert REPORT.md into a Professional Microsoft Word (.docx) Document
Applies publication-grade typography, formatted tables, code styling, and clean document hierarchy.
"""
import os
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

REPORT_MD_PATH = Path("REPORT.md")
OUTPUT_DOCX_PATH = Path("REPORT.docx")

# Styling Palette (Executive Navy & Slate Theme)
COLOR_PRIMARY = RGBColor(30, 58, 138)      # #1E3A8A Navy
COLOR_SECONDARY = RGBColor(71, 85, 105)    # #475569 Slate
COLOR_TEXT = RGBColor(34, 34, 34)          # #222222 Near Black
COLOR_MUTED = RGBColor(100, 116, 139)      # #64748B Muted Gray
COLOR_HEADER_BG = "1E3A8A"                 # Hex for XML
COLOR_ZEBRA_BG = "F8FAFC"                  # Hex for XML
COLOR_BORDER = "CBD5E1"                    # Hex for XML
COLOR_CODE_BG = "F1F5F9"                   # Hex for XML


def set_cell_background(cell, fill_hex):
    """Set background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """Set padding inside a table cell (in dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def set_table_borders(table):
    """Apply elegant thin borders to table."""
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{COLOR_HEADER_BG}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(borders)


def add_formatted_runs(paragraph, text, base_font_size=10.5, is_italic=False, color=COLOR_TEXT):
    """Parse inline Markdown (**bold**, *italic*, `code`, [link](url)) into Word runs."""
    tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))', text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.name = "Calibri"
        run.font.size = Pt(base_font_size)
        run.font.color.rgb = color
        run.italic = is_italic

        if token.startswith('**') and token.endswith('**'):
            run.text = token[2:-2]
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run.text = token[1:-1]
            run.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run.text = token[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(base_font_size - 0.5)
            run.font.color.rgb = RGBColor(180, 40, 40)
        elif token.startswith('[') and token.endswith(')'):
            m = re.match(r'\[(.*?)\]\((.*?)\)', token)
            if m:
                link_text, link_url = m.groups()
                run.text = link_text
                run.font.color.rgb = RGBColor(27, 77, 186)
                run.underline = True
            else:
                run.text = token
        else:
            run.text = token


def convert_markdown_to_docx(md_path: Path, docx_path: Path):
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    # Configure Margins: 1 inch (72 pt) all around
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Normal Style setup
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = COLOR_TEXT
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(4)

    i = 0
    n = len(lines)

    while i < n:
        raw_line = lines[i]
        line = raw_line.rstrip()

        # 1. Blank Line
        if not line.strip():
            i += 1
            continue

        # 1.5 Markdown Images (![Caption](image_path))
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)', line.strip())
        if img_match:
            caption, img_rel_path = img_match.groups()
            clean_path = img_rel_path.replace("file:///", "").replace("file://", "").strip()
            img_file = Path(clean_path)
            if not img_file.is_absolute():
                img_file = (md_path.parent / img_file).resolve()
            
            if img_file.exists() and img_file.is_file():
                try:
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(10)
                    p_img.paragraph_format.space_after = Pt(2)
                    run_img = p_img.add_run()
                    run_img.add_picture(str(img_file), width=Inches(6.0))

                    if caption:
                        p_cap = doc.add_paragraph()
                        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_cap.paragraph_format.space_before = Pt(1)
                        p_cap.paragraph_format.space_after = Pt(10)
                        run_cap = p_cap.add_run(f"Figure: {caption}")
                        run_cap.font.name = "Calibri"
                        run_cap.font.size = Pt(8.5)
                        run_cap.italic = True
                        run_cap.font.color.rgb = COLOR_MUTED
                except Exception as err:
                    print(f"Warning: could not add picture {img_file}: {err}")
            else:
                print(f"Warning: image file not found: {img_file}")
            i += 1
            continue

        # 2. Horizontal Divider
        if line.strip() in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="{COLOR_BORDER}"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue

        # 3. Document Title (# Title)
        if line.startswith("# ") and not line.startswith("## "):
            title_text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(title_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.bold = True
            run.font.color.rgb = COLOR_PRIMARY
            i += 1
            continue

        # 4. Document Subtitle (## Subtitle right after Title)
        if line.startswith("## Comprehensive Technical Project Report"):
            subtitle_text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(14)
            run = p.add_run(subtitle_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.italic = True
            run.font.color.rgb = COLOR_SECONDARY
            i += 1
            continue

        # 5. Heading 1 (## Section)
        if line.startswith("## "):
            h_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(15)
            run.bold = True
            run.font.color.rgb = COLOR_PRIMARY
            i += 1
            continue

        # 6. Heading 2 (### Sub-section)
        if line.startswith("### "):
            h_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(12.5)
            run.bold = True
            run.font.color.rgb = COLOR_PRIMARY
            i += 1
            continue

        # 7. Heading 3 (#### Sub-sub-section)
        if line.startswith("#### "):
            h_text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = COLOR_SECONDARY
            i += 1
            continue

        # 8. Code Blocks or ASCII Diagrams (```)
        if line.startswith("```"):
            code_lines = []
            i += 1
            while i < n and not lines[i].rstrip().startswith("```"):
                code_lines.append(lines[i].rstrip('\n'))
                i += 1
            i += 1  # skip closing ```

            # Render code box as a single-cell table with background
            code_text = "\n".join(code_lines)
            code_table = doc.add_table(rows=1, cols=1)
            code_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            code_cell = code_table.cell(0, 0)
            set_cell_background(code_cell, COLOR_CODE_BG)
            set_cell_margins(code_cell, top=100, bottom=100, left=150, right=150)
            
            p = code_cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(code_text)
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(30, 41, 59)
            
            # Spacer after table
            sp = doc.add_paragraph()
            sp.paragraph_format.space_before = Pt(0)
            sp.paragraph_format.space_after = Pt(4)
            continue

        # 9. LaTeX Math Equations ($$...$$)
        if line.startswith("$$") or (line.strip().startswith("$$") and line.strip().endswith("$$")):
            math_lines = []
            if line.strip() == "$$":
                i += 1
                while i < n and lines[i].strip() != "$$":
                    math_lines.append(lines[i].strip())
                    i += 1
                i += 1
            else:
                m_content = line.strip().strip("$").strip()
                math_lines.append(m_content)
                i += 1

            math_text = "   ".join(math_lines)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(math_text)
            run.font.name = "Cambria Math"
            run.font.size = Pt(10.5)
            run.italic = True
            run.font.color.rgb = COLOR_PRIMARY
            continue

        # 10. Markdown Tables (| ... |)
        if line.startswith("|") and line.endswith("|"):
            table_raw_lines = []
            while i < n and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_raw_lines.append(lines[i].strip())
                i += 1

            if len(table_raw_lines) >= 2:
                # Parse headers
                headers = [c.strip() for c in table_raw_lines[0].strip('|').split('|')]
                # Skip row 1 if it's separator (| :--- | :---: |)
                start_row = 2 if re.match(r'^[|\s:-]+$', table_raw_lines[1]) else 1

                data_rows = []
                for tr in table_raw_lines[start_row:]:
                    row_cells = [c.strip() for c in tr.strip('|').split('|')]
                    # Pad if cells count doesn't match
                    while len(row_cells) < len(headers):
                        row_cells.append("")
                    data_rows.append(row_cells[:len(headers)])

                word_table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
                word_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(word_table)

                # Style Header Row
                for col_idx, h_text in enumerate(headers):
                    cell = word_table.cell(0, col_idx)
                    set_cell_background(cell, COLOR_HEADER_BG)
                    set_cell_margins(cell, top=140, bottom=140, left=120, right=120)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    run = p.add_run(h_text.replace('**', ''))
                    run.font.name = "Calibri"
                    run.font.size = Pt(9.5)
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)

                # Style Data Rows
                for row_idx, r_data in enumerate(data_rows):
                    is_even = (row_idx % 2 == 1)
                    for col_idx, val in enumerate(r_data):
                        cell = word_table.cell(row_idx + 1, col_idx)
                        if is_even:
                            set_cell_background(cell, COLOR_ZEBRA_BG)
                        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(0)
                        p.paragraph_format.space_after = Pt(0)
                        add_formatted_runs(p, val, base_font_size=9.0)

                # Spacer after table
                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(2)
                sp.paragraph_format.space_after = Pt(4)
                continue

        # 11. Unordered Lists (- or *)
        if re.match(r'^[-*]\s+', line):
            bullet_text = re.sub(r'^[-*]\s+', '', line)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_runs(p, bullet_text, base_font_size=10.0)
            i += 1
            continue

        # 12. Ordered Lists (1., 2., etc.)
        if re.match(r'^\d+\.\s+', line):
            num_match = re.match(r'^(\d+\.)\s+(.*)', line)
            if num_match:
                prefix, num_text = num_match.groups()
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(2)
                p_bold = p.add_run(f"{prefix} ")
                p_bold.bold = True
                p_bold.font.color.rgb = COLOR_PRIMARY
                add_formatted_runs(p, num_text, base_font_size=10.0)
            i += 1
            continue

        # 13. Blockquotes (> ...)
        if line.startswith("> "):
            quote_text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="8" w:color="{COLOR_HEADER_BG}"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            add_formatted_runs(p, quote_text, base_font_size=10.0, is_italic=True, color=COLOR_SECONDARY)
            i += 1
            continue

        # 14. Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        add_formatted_runs(p, line, base_font_size=10.5)
        i += 1

    doc.save(str(docx_path))
    print(f"Successfully generated Word document: {docx_path.resolve()}")


if __name__ == "__main__":
    convert_markdown_to_docx(REPORT_MD_PATH, OUTPUT_DOCX_PATH)
